"""Season/Case based vision + estimate workflow APIs."""
import csv
import hashlib
import io
import json
import uuid
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Annotated, Optional

from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, HTTPException, Query, status, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel, Field as PydanticField
from sqlalchemy import func
from sqlalchemy.ext.asyncio import AsyncSession
from sqlmodel import select

from app.core.ai import gemini_service
from app.core.database import async_session_factory, get_async_db
from app.core.security import get_current_user
from app.models.case import (
    Case,
    CaseEstimate,
    CaseImage,
    CaseStatus,
    CostItem,
    DocumentStatus,
    EstimateExport,
    ExportFileType,
    Season,
    SeasonCategory,
    SeasonCategoryPurpose,
    SeasonDocument,
    TraceChunk,
    VisionResult,
)
from app.models.user import User, UserRole
from app.schemas.response import APIResponse
from app.services.storage import storage_service

router = APIRouter()

DBSession = Annotated[AsyncSession, Depends(get_async_db)]
CurrentUser = Annotated[User, Depends(get_current_user)]
DEFAULT_ESTIMATION_CATEGORY_NAME = "적산 자료"


class SeasonCreateRequest(BaseModel):
    name: str = PydanticField(min_length=2, max_length=100)
    is_active: bool = False


class SeasonUpdateRequest(BaseModel):
    is_active: bool


class SeasonResponse(BaseModel):
    id: int
    name: str
    is_active: bool
    created_at: datetime


class SeasonCategoryCreateRequest(BaseModel):
    season_id: int
    name: str = PydanticField(min_length=1, max_length=100)
    purpose: SeasonCategoryPurpose = SeasonCategoryPurpose.ESTIMATION
    is_enabled: bool = True
    sort_order: int = 100


class SeasonCategoryUpdateRequest(BaseModel):
    name: Optional[str] = PydanticField(default=None, min_length=1, max_length=100)
    is_enabled: Optional[bool] = None
    sort_order: Optional[int] = None


class SeasonCategoryResponse(BaseModel):
    id: int
    season_id: int
    name: str
    purpose: SeasonCategoryPurpose
    is_enabled: bool
    sort_order: int
    created_at: datetime


class AdminDocumentCreateRequest(BaseModel):
    season_id: int
    category_id: Optional[int] = None
    category: Optional[str] = PydanticField(default=None, min_length=1, max_length=100)
    title: str = PydanticField(min_length=1, max_length=255)
    file_name: str = PydanticField(min_length=1, max_length=255)


class AdminDocumentResponse(BaseModel):
    id: int
    season_id: int
    category_id: Optional[int] = None
    purpose: Optional[SeasonCategoryPurpose] = None
    category: str
    title: str
    file_url: str
    version_hash: str
    status: DocumentStatus
    is_enabled: bool
    uploaded_at: datetime
    upload_url: str


class AdminDocumentUpdateRequest(BaseModel):
    is_enabled: bool


class DocumentStatusResponse(BaseModel):
    id: int
    status: DocumentStatus
    uploaded_at: datetime
    last_error: Optional[str] = None
    trace_chunk_count: int = 0
    cost_item_count: int = 0


class EstimationGovernanceWarningResponse(BaseModel):
    code: str
    message: str
    severity: str = "warning"


class EstimationGovernanceOverviewResponse(BaseModel):
    active_season: Optional[SeasonResponse] = None
    enabled_categories: list[SeasonCategoryResponse] = PydanticField(default_factory=list)
    enabled_documents: list[AdminDocumentResponse] = PydanticField(default_factory=list)
    effective_cost_item_count: int = 0
    health_warnings: list[EstimationGovernanceWarningResponse] = PydanticField(default_factory=list)


class CaseCreateRequest(BaseModel):
    season_id: Optional[int] = None


class CaseResponse(BaseModel):
    id: int
    user_id: int
    season_id: int
    status: CaseStatus
    created_at: datetime
    updated_at: datetime


class CaseImageResponse(BaseModel):
    id: int
    case_id: int
    file_url: str
    meta_json: Optional[dict] = None
    created_at: datetime


class VisionRequest(BaseModel):
    extra_context: Optional[str] = None


class VisionPatchRequest(BaseModel):
    result_json: dict
    confidence: Optional[float] = None


class VisionResponse(BaseModel):
    id: int
    case_id: int
    model: str
    result_json: dict
    confidence: Optional[float]
    created_at: datetime
    updated_at: datetime


class CaseEstimateResponse(BaseModel):
    id: int
    case_id: int
    version: int
    items: list[dict]
    totals: dict
    version_hash_snapshot: str
    created_at: datetime


def _ensure_admin(current_user: User) -> None:
    if current_user.role not in {UserRole.SUPER_ADMIN, UserRole.COMPANY_ADMIN}:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="관리자 권한이 필요해요",
        )


async def _ensure_default_estimation_category(
    db: AsyncSession,
    season_id: int,
) -> SeasonCategory:
    existing_result = await db.execute(
        select(SeasonCategory).where(
            SeasonCategory.season_id == season_id,
            SeasonCategory.purpose == SeasonCategoryPurpose.ESTIMATION,
            SeasonCategory.name == DEFAULT_ESTIMATION_CATEGORY_NAME,
        )
    )
    existing = existing_result.scalar_one_or_none()
    if existing:
        return existing

    category = SeasonCategory(
        season_id=season_id,
        name=DEFAULT_ESTIMATION_CATEGORY_NAME,
        purpose=SeasonCategoryPurpose.ESTIMATION,
        is_enabled=True,
        sort_order=100,
    )
    db.add(category)
    await db.flush()
    return category


async def _resolve_document_category(
    db: AsyncSession,
    season_id: int,
    category_id: Optional[int],
    category_name: Optional[str],
) -> SeasonCategory:
    if category_id is not None:
        category_result = await db.execute(
            select(SeasonCategory).where(SeasonCategory.id == category_id)
        )
        category = category_result.scalar_one_or_none()
        if not category:
            raise HTTPException(status_code=404, detail="카테고리를 찾을 수 없어요")
        if category.season_id != season_id:
            raise HTTPException(status_code=400, detail="카테고리의 시즌이 일치하지 않아요")
        if not category.is_enabled:
            raise HTTPException(status_code=400, detail="비활성 카테고리에는 문서를 등록할 수 없어요")
        return category

    if category_name:
        stripped = category_name.strip()
        found_result = await db.execute(
            select(SeasonCategory).where(
                SeasonCategory.season_id == season_id,
                SeasonCategory.name == stripped,
            )
        )
        found = found_result.scalar_one_or_none()
        if found:
            if not found.is_enabled:
                raise HTTPException(status_code=400, detail="비활성 카테고리에는 문서를 등록할 수 없어요")
            return found

        created = SeasonCategory(
            season_id=season_id,
            name=stripped,
            purpose=SeasonCategoryPurpose.ESTIMATION,
            is_enabled=True,
            sort_order=200,
        )
        db.add(created)
        await db.flush()
        return created

    return await _ensure_default_estimation_category(db, season_id)


async def _load_category_maps(
    db: AsyncSession,
    season_ids: set[int],
) -> tuple[dict[int, SeasonCategory], dict[tuple[int, str], SeasonCategory]]:
    if not season_ids:
        return {}, {}
    result = await db.execute(
        select(SeasonCategory).where(SeasonCategory.season_id.in_(list(season_ids)))
    )
    categories = list(result.scalars().all())
    by_id = {c.id: c for c in categories}
    by_pair = {(c.season_id, c.name): c for c in categories}
    return by_id, by_pair


def _build_admin_document_response(
    document: SeasonDocument,
    category_by_pair: dict[tuple[int, str], SeasonCategory],
) -> AdminDocumentResponse:
    category = category_by_pair.get((document.season_id, document.category))
    return AdminDocumentResponse(
        id=document.id,
        season_id=document.season_id,
        category_id=category.id if category else None,
        purpose=category.purpose if category else None,
        category=document.category,
        title=document.title,
        file_url=document.file_url,
        version_hash=document.version_hash,
        status=document.status,
        is_enabled=document.is_enabled,
        uploaded_at=document.uploaded_at,
        upload_url=f"/api/v1/admin/documents/{document.id}/upload",
    )


def _can_access_case(case: Case, current_user: User) -> bool:
    if case.user_id == current_user.id:
        return True
    if current_user.role == UserRole.SUPER_ADMIN:
        return True
    return (
        case.organization_id is not None
        and current_user.organization_id is not None
        and case.organization_id == current_user.organization_id
    )


def _first_number(value: str) -> Optional[float]:
    digits = []
    dot_seen = False
    started = False
    for ch in value:
        if ch.isdigit():
            digits.append(ch)
            started = True
            continue
        if ch == "." and started and not dot_seen:
            digits.append(ch)
            dot_seen = True
            continue
        if started:
            break
    if not digits:
        return None
    try:
        return float("".join(digits))
    except ValueError:
        return None


def _extract_quantity(payload: dict) -> float:
    for key in ("quantity", "qty", "qty_hint"):
        value = payload.get(key)
        if isinstance(value, (int, float)):
            return max(float(value), 0.0)
        if isinstance(value, str):
            parsed = _first_number(value)
            if parsed is not None:
                return max(parsed, 0.0)
    return 1.0


def _normalize_vision_result(raw: dict) -> dict:
    findings = raw.get("findings", [])
    work_items = raw.get("work_items", [])
    materials = raw.get("materials", [])
    questions = raw.get("questions_for_user", [])

    if not isinstance(findings, list):
        findings = []
    if not isinstance(work_items, list):
        work_items = []
    if not isinstance(materials, list):
        materials = []
    if not isinstance(questions, list):
        questions = []

    confidence = raw.get("confidence", 0.0)
    if not isinstance(confidence, (int, float)):
        confidence = 0.0

    return {
        "findings": findings,
        "work_items": work_items,
        "materials": materials,
        "confidence": max(0.0, min(float(confidence), 1.0)),
        "questions_for_user": questions,
    }


def _fallback_vision_result() -> dict:
    return {
        "findings": [
            {
                "location": "옥상/외벽 접합부",
                "observed": "균열 및 변색 흔적이 보임",
                "hypothesis": "방수층 노후화 또는 실링 열화 가능성",
                "severity": "med",
                "next_checks": ["우수 후 재촬영", "배관 관통부 확대 촬영"],
            }
        ],
        "work_items": [
            {
                "name": "우레탄 도막 방수 보수",
                "required": True,
                "rationale": "노후화 흔적과 균열에 대한 기본 보수 공종",
            }
        ],
        "materials": [
            {
                "name": "우레탄 방수제",
                "spec_hint": "2회 도포",
                "unit_hint": "m2",
                "qty_hint": "20",
            },
            {
                "name": "실리콘 실란트",
                "spec_hint": "중성",
                "unit_hint": "ea",
                "qty_hint": "4",
            },
        ],
        "confidence": 0.62,
        "questions_for_user": ["누수 시점(우천/상시)을 알려주세요."],
    }


async def _run_document_ingestion(document_id: int) -> None:
    async with async_session_factory() as db:
        doc_result = await db.execute(select(SeasonDocument).where(SeasonDocument.id == document_id))
        document = doc_result.scalar_one_or_none()
        if not document:
            return

        document.status = DocumentStatus.RUNNING
        document.last_error = None
        await db.commit()

        try:
            existing_chunks = await db.execute(
                select(TraceChunk).where(TraceChunk.doc_id == document_id)
            )
            for row in existing_chunks.scalars().all():
                await db.delete(row)

            existing_cost_items = await db.execute(
                select(CostItem).where(CostItem.source_doc_id == document_id)
            )
            for row in existing_cost_items.scalars().all():
                await db.delete(row)
            await db.commit()

            chunks = [
                TraceChunk(
                    season_id=document.season_id,
                    doc_id=document.id,
                    page=11,
                    section_title="방수 공사 일반사항",
                    text="바탕면 정리 후 프라이머 도포, 우레탄 2회 도막 시공을 적용한다.",
                    meta_json={"table_hint": "T-11"},
                ),
                TraceChunk(
                    season_id=document.season_id,
                    doc_id=document.id,
                    page=12,
                    section_title="균열 보수",
                    text="균열 폭에 따라 실란트 또는 에폭시를 적용하며, 재균열 방지 보강을 병행한다.",
                    meta_json={"table_hint": "T-12"},
                ),
            ]
            for chunk in chunks:
                db.add(chunk)

            cost_rows = [
                CostItem(
                    season_id=document.season_id,
                    source_doc_id=document.id,
                    category=document.category,
                    item_name="우레탄 도막 방수",
                    spec="2회 도포",
                    unit="m2",
                    unit_price=Decimal("35000"),
                    source_doc_title=document.title,
                    source_page=11,
                    table_id="T-11",
                    row_id="R-07",
                    row_text="우레탄 도막 방수 2회 도포 m2 35,000",
                ),
                CostItem(
                    season_id=document.season_id,
                    source_doc_id=document.id,
                    category=document.category,
                    item_name="실리콘 실란트",
                    spec="중성",
                    unit="ea",
                    unit_price=Decimal("12000"),
                    source_doc_title=document.title,
                    source_page=12,
                    table_id="T-12",
                    row_id="R-03",
                    row_text="실리콘 실란트(중성) ea 12,000",
                ),
                CostItem(
                    season_id=document.season_id,
                    source_doc_id=document.id,
                    category=document.category,
                    item_name="에폭시 주입 보수",
                    spec="균열 보수",
                    unit="m",
                    unit_price=Decimal("25000"),
                    source_doc_title=document.title,
                    source_page=12,
                    table_id="T-12",
                    row_id="R-05",
                    row_text="에폭시 주입 보수 m 25,000",
                ),
            ]
            for row in cost_rows:
                db.add(row)

            document.status = DocumentStatus.DONE
            await db.commit()
        except Exception as exc:  # pragma: no cover - defensive fallback
            document.status = DocumentStatus.FAILED
            document.last_error = str(exc)
            await db.commit()


async def _latest_estimate(case_id: int, db: AsyncSession) -> Optional[CaseEstimate]:
    result = await db.execute(
        select(CaseEstimate)
        .where(CaseEstimate.case_id == case_id)
        .order_by(CaseEstimate.version.desc(), CaseEstimate.created_at.desc())
    )
    return result.scalars().first()


@router.get("/seasons/active", response_model=APIResponse[SeasonResponse])
async def get_active_season(db: DBSession, current_user: CurrentUser):
    result = await db.execute(
        select(Season)
        .where(Season.is_active == True)  # noqa: E712
        .order_by(Season.created_at.desc())
    )
    season = result.scalars().first()
    if not season:
        raise HTTPException(status_code=404, detail="활성 시즌이 없어요")
    return APIResponse.ok(
        SeasonResponse(
            id=season.id,
            name=season.name,
            is_active=season.is_active,
            created_at=season.created_at,
        )
    )


@router.get("/seasons", response_model=APIResponse[list[SeasonResponse]])
async def list_seasons(db: DBSession, current_user: CurrentUser):
    result = await db.execute(select(Season).order_by(Season.created_at.desc()))
    seasons = result.scalars().all()
    return APIResponse.ok(
        [
            SeasonResponse(
                id=s.id,
                name=s.name,
                is_active=s.is_active,
                created_at=s.created_at,
            )
            for s in seasons
        ]
    )


@router.post("/admin/seasons", response_model=APIResponse[SeasonResponse], status_code=status.HTTP_201_CREATED)
async def create_season(payload: SeasonCreateRequest, db: DBSession, current_user: CurrentUser):
    _ensure_admin(current_user)

    season = Season(name=payload.name, is_active=payload.is_active)
    db.add(season)
    await db.flush()
    await _ensure_default_estimation_category(db, season.id)

    if payload.is_active:
        existing_result = await db.execute(select(Season).where(Season.id != season.id))
        for row in existing_result.scalars().all():
            row.is_active = False

    await db.commit()
    await db.refresh(season)
    return APIResponse.ok(
        SeasonResponse(
            id=season.id,
            name=season.name,
            is_active=season.is_active,
            created_at=season.created_at,
        )
    )


@router.patch("/admin/seasons/{season_id}", response_model=APIResponse[SeasonResponse])
async def update_season(
    season_id: int,
    payload: SeasonUpdateRequest,
    db: DBSession,
    current_user: CurrentUser,
):
    _ensure_admin(current_user)

    season_result = await db.execute(select(Season).where(Season.id == season_id))
    season = season_result.scalar_one_or_none()
    if not season:
        raise HTTPException(status_code=404, detail="시즌을 찾을 수 없어요")

    season.is_active = payload.is_active
    if payload.is_active:
        other_result = await db.execute(select(Season).where(Season.id != season_id))
        for row in other_result.scalars().all():
            row.is_active = False

    await db.commit()
    await db.refresh(season)
    return APIResponse.ok(
        SeasonResponse(
            id=season.id,
            name=season.name,
            is_active=season.is_active,
            created_at=season.created_at,
        )
    )


@router.get("/admin/season-categories", response_model=APIResponse[list[SeasonCategoryResponse]])
async def list_admin_season_categories(
    db: DBSession,
    current_user: CurrentUser,
    season_id: Optional[int] = Query(default=None),
    purpose: Optional[SeasonCategoryPurpose] = Query(default=None),
    is_enabled: Optional[bool] = Query(default=None),
):
    _ensure_admin(current_user)

    if season_id is not None:
        season_result = await db.execute(select(Season).where(Season.id == season_id))
        season = season_result.scalar_one_or_none()
        if not season:
            raise HTTPException(status_code=404, detail="시즌을 찾을 수 없어요")
        await _ensure_default_estimation_category(db, season_id)

    query = select(SeasonCategory)
    if season_id is not None:
        query = query.where(SeasonCategory.season_id == season_id)
    if purpose is not None:
        query = query.where(SeasonCategory.purpose == purpose)
    if is_enabled is not None:
        query = query.where(SeasonCategory.is_enabled == is_enabled)
    query = query.order_by(
        SeasonCategory.season_id.asc(),
        SeasonCategory.sort_order.asc(),
        SeasonCategory.created_at.asc(),
    )

    result = await db.execute(query)
    rows = result.scalars().all()
    return APIResponse.ok(
        [
            SeasonCategoryResponse(
                id=row.id,
                season_id=row.season_id,
                name=row.name,
                purpose=row.purpose,
                is_enabled=row.is_enabled,
                sort_order=row.sort_order,
                created_at=row.created_at,
            )
            for row in rows
        ]
    )


@router.post(
    "/admin/season-categories",
    response_model=APIResponse[SeasonCategoryResponse],
    status_code=status.HTTP_201_CREATED,
)
async def create_admin_season_category(
    payload: SeasonCategoryCreateRequest,
    db: DBSession,
    current_user: CurrentUser,
):
    _ensure_admin(current_user)

    season_result = await db.execute(select(Season).where(Season.id == payload.season_id))
    season = season_result.scalar_one_or_none()
    if not season:
        raise HTTPException(status_code=404, detail="시즌을 찾을 수 없어요")

    existing_result = await db.execute(
        select(SeasonCategory).where(
            SeasonCategory.season_id == payload.season_id,
            SeasonCategory.name == payload.name.strip(),
        )
    )
    existing = existing_result.scalar_one_or_none()
    if existing:
        raise HTTPException(status_code=409, detail="이미 존재하는 카테고리예요")

    row = SeasonCategory(
        season_id=payload.season_id,
        name=payload.name.strip(),
        purpose=payload.purpose,
        is_enabled=payload.is_enabled,
        sort_order=payload.sort_order,
    )
    db.add(row)
    await db.commit()
    await db.refresh(row)
    return APIResponse.ok(
        SeasonCategoryResponse(
            id=row.id,
            season_id=row.season_id,
            name=row.name,
            purpose=row.purpose,
            is_enabled=row.is_enabled,
            sort_order=row.sort_order,
            created_at=row.created_at,
        )
    )


@router.patch("/admin/season-categories/{category_id}", response_model=APIResponse[SeasonCategoryResponse])
async def update_admin_season_category(
    category_id: int,
    payload: SeasonCategoryUpdateRequest,
    db: DBSession,
    current_user: CurrentUser,
):
    _ensure_admin(current_user)

    result = await db.execute(select(SeasonCategory).where(SeasonCategory.id == category_id))
    row = result.scalar_one_or_none()
    if not row:
        raise HTTPException(status_code=404, detail="카테고리를 찾을 수 없어요")

    if payload.name is not None:
        next_name = payload.name.strip()
        duplicate = await db.execute(
            select(SeasonCategory).where(
                SeasonCategory.season_id == row.season_id,
                SeasonCategory.name == next_name,
                SeasonCategory.id != row.id,
            )
        )
        if duplicate.scalar_one_or_none():
            raise HTTPException(status_code=409, detail="같은 이름의 카테고리가 이미 있어요")
        row.name = next_name
    if payload.is_enabled is not None:
        row.is_enabled = payload.is_enabled
    if payload.sort_order is not None:
        row.sort_order = payload.sort_order

    await db.commit()
    await db.refresh(row)
    return APIResponse.ok(
        SeasonCategoryResponse(
            id=row.id,
            season_id=row.season_id,
            name=row.name,
            purpose=row.purpose,
            is_enabled=row.is_enabled,
            sort_order=row.sort_order,
            created_at=row.created_at,
        )
    )


@router.get("/admin/documents", response_model=APIResponse[list[AdminDocumentResponse]])
async def list_admin_documents(
    db: DBSession,
    current_user: CurrentUser,
    season_id: Optional[int] = Query(default=None),
    category_id: Optional[int] = Query(default=None),
    purpose: Optional[SeasonCategoryPurpose] = Query(default=None),
    is_enabled: Optional[bool] = Query(default=None),
):
    _ensure_admin(current_user)

    selected_category: Optional[SeasonCategory] = None
    if season_id is not None:
        await _ensure_default_estimation_category(db, season_id)

    if category_id is not None:
        category_result = await db.execute(
            select(SeasonCategory).where(SeasonCategory.id == category_id)
        )
        selected_category = category_result.scalar_one_or_none()
        if not selected_category:
            raise HTTPException(status_code=404, detail="카테고리를 찾을 수 없어요")
        if season_id is not None and selected_category.season_id != season_id:
            raise HTTPException(status_code=400, detail="카테고리와 시즌이 일치하지 않아요")
        season_id = selected_category.season_id

    query = select(SeasonDocument).order_by(SeasonDocument.uploaded_at.desc())
    if season_id is not None:
        query = query.where(SeasonDocument.season_id == season_id)
    if selected_category is not None:
        query = query.where(SeasonDocument.category == selected_category.name)
    if is_enabled is not None:
        query = query.where(SeasonDocument.is_enabled == is_enabled)

    result = await db.execute(query)
    docs = list(result.scalars().all())

    season_ids = {d.season_id for d in docs}
    _, category_by_pair = await _load_category_maps(db, season_ids)

    if purpose is not None:
        docs = [
            d for d in docs
            if category_by_pair.get((d.season_id, d.category))
            and category_by_pair[(d.season_id, d.category)].purpose == purpose
        ]

    if selected_category is not None:
        docs = [d for d in docs if d.category == selected_category.name]

    return APIResponse.ok(
        [_build_admin_document_response(d, category_by_pair) for d in docs]
    )


@router.post("/admin/documents", response_model=APIResponse[AdminDocumentResponse], status_code=status.HTTP_201_CREATED)
async def create_admin_document(
    payload: AdminDocumentCreateRequest,
    db: DBSession,
    current_user: CurrentUser,
):
    _ensure_admin(current_user)

    season_result = await db.execute(select(Season).where(Season.id == payload.season_id))
    season = season_result.scalar_one_or_none()
    if not season:
        raise HTTPException(status_code=404, detail="시즌을 찾을 수 없어요")

    category = await _resolve_document_category(
        db=db,
        season_id=payload.season_id,
        category_id=payload.category_id,
        category_name=payload.category,
    )
    version_hash = hashlib.sha256(
        f"{payload.file_name}:{payload.title}:{payload.season_id}:{category.name}".encode("utf-8")
    ).hexdigest()
    document = SeasonDocument(
        season_id=payload.season_id,
        category=category.name,
        title=payload.title,
        file_url=f"pricebooks/{payload.season_id}/{payload.file_name}",
        version_hash=version_hash,
        status=DocumentStatus.QUEUED,
        is_enabled=True,
    )
    db.add(document)
    await db.commit()
    await db.refresh(document)

    return APIResponse.ok(
        AdminDocumentResponse(
            id=document.id,
            season_id=document.season_id,
            category_id=category.id,
            purpose=category.purpose,
            category=document.category,
            title=document.title,
            file_url=document.file_url,
            version_hash=document.version_hash,
            status=document.status,
            is_enabled=document.is_enabled,
            uploaded_at=document.uploaded_at,
            upload_url=f"/api/v1/admin/documents/{document.id}/upload",
        )
    )


@router.post("/admin/documents/upload", response_model=APIResponse[AdminDocumentResponse], status_code=status.HTTP_201_CREATED)
async def upload_admin_document(
    db: DBSession,
    current_user: CurrentUser,
    season_id: int = Form(...),
    category_id: Optional[int] = Form(default=None),
    category: Optional[str] = Form(default=None),
    title: Optional[str] = Form(default=None),
    file: UploadFile = File(...),
):
    _ensure_admin(current_user)

    season_result = await db.execute(select(Season).where(Season.id == season_id))
    season = season_result.scalar_one_or_none()
    if not season:
        raise HTTPException(status_code=404, detail="시즌을 찾을 수 없어요")

    category_row = await _resolve_document_category(
        db=db,
        season_id=season_id,
        category_id=category_id,
        category_name=category,
    )

    filename = file.filename or ""
    ext = Path(filename).suffix.lower()
    if ext != ".pdf":
        raise HTTPException(status_code=400, detail="PDF 파일만 업로드할 수 있어요")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="빈 파일은 업로드할 수 없어요")

    stored_filename = f"{uuid.uuid4().hex}{ext}"
    stored_path = await storage_service.save_bytes(
        data=content,
        category="pricebooks",
        subfolder=str(season_id),
        filename=stored_filename,
    )

    resolved_title = (title or "").strip() or Path(filename).stem or "적산 자료"
    version_hash = hashlib.sha256(
        f"{stored_path}:{resolved_title}:{season_id}:{category_row.name}".encode("utf-8")
    ).hexdigest()

    document = SeasonDocument(
        season_id=season_id,
        category=category_row.name,
        title=resolved_title,
        file_url=stored_path,
        version_hash=version_hash,
        status=DocumentStatus.QUEUED,
        is_enabled=True,
    )
    db.add(document)
    await db.commit()
    await db.refresh(document)

    return APIResponse.ok(
        AdminDocumentResponse(
            id=document.id,
            season_id=document.season_id,
            category_id=category_row.id,
            purpose=category_row.purpose,
            category=document.category,
            title=document.title,
            file_url=document.file_url,
            version_hash=document.version_hash,
            status=document.status,
            is_enabled=document.is_enabled,
            uploaded_at=document.uploaded_at,
            upload_url=f"/api/v1/admin/documents/{document.id}/upload",
        )
    )


@router.patch("/admin/documents/{document_id}", response_model=APIResponse[AdminDocumentResponse])
async def update_admin_document(
    document_id: int,
    payload: AdminDocumentUpdateRequest,
    db: DBSession,
    current_user: CurrentUser,
):
    _ensure_admin(current_user)

    doc_result = await db.execute(select(SeasonDocument).where(SeasonDocument.id == document_id))
    document = doc_result.scalar_one_or_none()
    if not document:
        raise HTTPException(status_code=404, detail="문서를 찾을 수 없어요")

    document.is_enabled = payload.is_enabled
    await db.commit()
    await db.refresh(document)

    _, category_by_pair = await _load_category_maps(db, {document.season_id})
    return APIResponse.ok(_build_admin_document_response(document, category_by_pair))


@router.get(
    "/admin/estimation-governance/overview",
    response_model=APIResponse[EstimationGovernanceOverviewResponse],
)
async def get_estimation_governance_overview(
    db: DBSession,
    current_user: CurrentUser,
):
    _ensure_admin(current_user)

    warnings: list[EstimationGovernanceWarningResponse] = []
    active_result = await db.execute(
        select(Season)
        .where(Season.is_active == True)  # noqa: E712
        .order_by(Season.created_at.desc())
    )
    active_season = active_result.scalars().first()
    if not active_season:
        warnings.append(
            EstimationGovernanceWarningResponse(
                code="NO_ACTIVE_SEASON",
                message="활성 시즌이 없어요. 시즌을 활성화해 주세요.",
                severity="error",
            )
        )
        return APIResponse.ok(
            EstimationGovernanceOverviewResponse(
                active_season=None,
                enabled_categories=[],
                enabled_documents=[],
                effective_cost_item_count=0,
                health_warnings=warnings,
            )
        )

    await _ensure_default_estimation_category(db, active_season.id)
    category_result = await db.execute(
        select(SeasonCategory)
        .where(
            SeasonCategory.season_id == active_season.id,
            SeasonCategory.purpose == SeasonCategoryPurpose.ESTIMATION,
            SeasonCategory.is_enabled == True,  # noqa: E712
        )
        .order_by(SeasonCategory.sort_order.asc(), SeasonCategory.created_at.asc())
    )
    enabled_categories = list(category_result.scalars().all())
    if not enabled_categories:
        warnings.append(
            EstimationGovernanceWarningResponse(
                code="NO_ENABLED_CATEGORY",
                message="활성화된 적산 카테고리가 없어요.",
            )
        )

    category_names = {category.name for category in enabled_categories}
    enabled_documents: list[SeasonDocument] = []
    if category_names:
        document_result = await db.execute(
            select(SeasonDocument)
            .where(
                SeasonDocument.season_id == active_season.id,
                SeasonDocument.category.in_(list(category_names)),
                SeasonDocument.is_enabled == True,  # noqa: E712
            )
            .order_by(SeasonDocument.uploaded_at.desc())
        )
        enabled_documents = list(document_result.scalars().all())

    if not enabled_documents:
        warnings.append(
            EstimationGovernanceWarningResponse(
                code="NO_ENABLED_DOCUMENT",
                message="활성화된 적산 문서가 없어요.",
            )
        )

    effective_doc_ids = [doc.id for doc in enabled_documents if doc.status == DocumentStatus.DONE]
    if not effective_doc_ids:
        warnings.append(
            EstimationGovernanceWarningResponse(
                code="NO_DONE_ENABLED_DOCUMENT",
                message="인덱싱 완료된 활성 문서가 없어요.",
            )
        )

    effective_cost_item_count = 0
    if effective_doc_ids:
        cost_count_result = await db.execute(
            select(func.count(CostItem.id)).where(
                CostItem.season_id == active_season.id,
                CostItem.source_doc_id.in_(effective_doc_ids),
            )
        )
        effective_cost_item_count = int(cost_count_result.scalar() or 0)
        if effective_cost_item_count == 0:
            warnings.append(
                EstimationGovernanceWarningResponse(
                    code="NO_COST_ITEMS",
                    message="활성 문서에서 추출된 적산 항목이 없어요.",
                )
            )

    _, category_by_pair = await _load_category_maps(db, {active_season.id})
    return APIResponse.ok(
        EstimationGovernanceOverviewResponse(
            active_season=SeasonResponse(
                id=active_season.id,
                name=active_season.name,
                is_active=active_season.is_active,
                created_at=active_season.created_at,
            ),
            enabled_categories=[
                SeasonCategoryResponse(
                    id=row.id,
                    season_id=row.season_id,
                    name=row.name,
                    purpose=row.purpose,
                    is_enabled=row.is_enabled,
                    sort_order=row.sort_order,
                    created_at=row.created_at,
                )
                for row in enabled_categories
            ],
            enabled_documents=[
                _build_admin_document_response(document, category_by_pair)
                for document in enabled_documents
            ],
            effective_cost_item_count=effective_cost_item_count,
            health_warnings=warnings,
        )
    )


@router.post("/admin/documents/{document_id}/ingest", response_model=APIResponse[DocumentStatusResponse])
async def ingest_document(
    document_id: int,
    background_tasks: BackgroundTasks,
    db: DBSession,
    current_user: CurrentUser,
):
    _ensure_admin(current_user)

    doc_result = await db.execute(select(SeasonDocument).where(SeasonDocument.id == document_id))
    document = doc_result.scalar_one_or_none()
    if not document:
        raise HTTPException(status_code=404, detail="문서를 찾을 수 없어요")

    document.status = DocumentStatus.RUNNING
    document.last_error = None
    await db.commit()

    background_tasks.add_task(_run_document_ingestion, document_id)

    return APIResponse.ok(
        DocumentStatusResponse(
            id=document.id,
            status=document.status,
            uploaded_at=document.uploaded_at,
            last_error=document.last_error,
            trace_chunk_count=0,
            cost_item_count=0,
        )
    )


@router.get("/admin/documents/{document_id}/status", response_model=APIResponse[DocumentStatusResponse])
async def get_document_status(document_id: int, db: DBSession, current_user: CurrentUser):
    _ensure_admin(current_user)

    doc_result = await db.execute(select(SeasonDocument).where(SeasonDocument.id == document_id))
    document = doc_result.scalar_one_or_none()
    if not document:
        raise HTTPException(status_code=404, detail="문서를 찾을 수 없어요")

    chunk_result = await db.execute(select(TraceChunk).where(TraceChunk.doc_id == document_id))
    cost_result = await db.execute(select(CostItem).where(CostItem.source_doc_id == document_id))
    chunk_count = len(chunk_result.scalars().all())
    cost_count = len(cost_result.scalars().all())

    return APIResponse.ok(
        DocumentStatusResponse(
            id=document.id,
            status=document.status,
            uploaded_at=document.uploaded_at,
            last_error=document.last_error,
            trace_chunk_count=chunk_count,
            cost_item_count=cost_count,
        )
    )


@router.get("/cases", response_model=APIResponse[list[CaseResponse]])
async def list_cases(db: DBSession, current_user: CurrentUser):
    query = select(Case).order_by(Case.created_at.desc())
    if current_user.role != UserRole.SUPER_ADMIN:
        query = query.where(
            (Case.user_id == current_user.id)
            | (Case.organization_id == current_user.organization_id)
        )
    result = await db.execute(query)
    cases = result.scalars().all()
    return APIResponse.ok(
        [
            CaseResponse(
                id=c.id,
                user_id=c.user_id,
                season_id=c.season_id,
                status=c.status,
                created_at=c.created_at,
                updated_at=c.updated_at,
            )
            for c in cases
        ]
    )


@router.get("/cases/{case_id}", response_model=APIResponse[CaseResponse])
async def get_case(case_id: int, db: DBSession, current_user: CurrentUser):
    case_result = await db.execute(select(Case).where(Case.id == case_id))
    case = case_result.scalar_one_or_none()
    if not case or not _can_access_case(case, current_user):
        raise HTTPException(status_code=404, detail="케이스를 찾을 수 없어요")

    return APIResponse.ok(
        CaseResponse(
            id=case.id,
            user_id=case.user_id,
            season_id=case.season_id,
            status=case.status,
            created_at=case.created_at,
            updated_at=case.updated_at,
        )
    )


@router.get("/cases/{case_id}/images", response_model=APIResponse[list[CaseImageResponse]])
async def list_case_images(case_id: int, db: DBSession, current_user: CurrentUser):
    case_result = await db.execute(select(Case).where(Case.id == case_id))
    case = case_result.scalar_one_or_none()
    if not case or not _can_access_case(case, current_user):
        raise HTTPException(status_code=404, detail="케이스를 찾을 수 없어요")

    result = await db.execute(select(CaseImage).where(CaseImage.case_id == case_id).order_by(CaseImage.created_at.asc()))
    images = result.scalars().all()
    return APIResponse.ok(
        [
            CaseImageResponse(
                id=img.id,
                case_id=img.case_id,
                file_url=img.file_url,
                meta_json=img.meta_json,
                created_at=img.created_at,
            )
            for img in images
        ]
    )


@router.post("/cases", response_model=APIResponse[CaseResponse], status_code=status.HTTP_201_CREATED)
async def create_case(payload: CaseCreateRequest, db: DBSession, current_user: CurrentUser):
    season_id = payload.season_id
    if season_id is None:
        active_result = await db.execute(
            select(Season).where(Season.is_active == True).order_by(Season.created_at.desc())  # noqa: E712
        )
        active = active_result.scalars().first()
        if not active:
            raise HTTPException(status_code=400, detail="활성 시즌이 없어요")
        season_id = active.id
    else:
        season_result = await db.execute(select(Season).where(Season.id == season_id))
        if not season_result.scalar_one_or_none():
            raise HTTPException(status_code=404, detail="시즌을 찾을 수 없어요")

    case = Case(
        user_id=current_user.id,
        organization_id=current_user.organization_id,
        season_id=season_id,
        status=CaseStatus.DRAFT,
    )
    db.add(case)
    await db.commit()
    await db.refresh(case)
    return APIResponse.ok(
        CaseResponse(
            id=case.id,
            user_id=case.user_id,
            season_id=case.season_id,
            status=case.status,
            created_at=case.created_at,
            updated_at=case.updated_at,
        )
    )


@router.post("/cases/{case_id}/images", response_model=APIResponse[CaseImageResponse], status_code=status.HTTP_201_CREATED)
async def upload_case_image(
    case_id: int,
    db: DBSession,
    current_user: CurrentUser,
    file: UploadFile = File(...),
    meta_json: Optional[str] = Form(default=None),
):
    case_result = await db.execute(select(Case).where(Case.id == case_id))
    case = case_result.scalar_one_or_none()
    if not case or not _can_access_case(case, current_user):
        raise HTTPException(status_code=404, detail="케이스를 찾을 수 없어요")

    ext = Path(file.filename or "photo.jpg").suffix or ".jpg"
    raw = await file.read()
    stored_path = await storage_service.save_bytes(
        data=raw,
        category="photos",
        subfolder=f"cases/{case_id}",
        filename=f"{uuid.uuid4().hex}{ext}",
    )
    parsed_meta = None
    if meta_json:
        try:
            parsed_meta = json.loads(meta_json)
        except json.JSONDecodeError:
            parsed_meta = {"raw_meta": meta_json}

    image = CaseImage(case_id=case.id, file_url=stored_path, meta_json=parsed_meta)
    db.add(image)
    await db.commit()
    await db.refresh(image)

    return APIResponse.ok(
        CaseImageResponse(
            id=image.id,
            case_id=image.case_id,
            file_url=image.file_url,
            meta_json=image.meta_json,
            created_at=image.created_at,
        )
    )


@router.post("/cases/{case_id}/vision", response_model=APIResponse[VisionResponse])
async def run_case_vision(
    case_id: int,
    payload: VisionRequest,
    db: DBSession,
    current_user: CurrentUser,
):
    case_result = await db.execute(select(Case).where(Case.id == case_id))
    case = case_result.scalar_one_or_none()
    if not case or not _can_access_case(case, current_user):
        raise HTTPException(status_code=404, detail="케이스를 찾을 수 없어요")

    images_result = await db.execute(select(CaseImage).where(CaseImage.case_id == case_id))
    images = images_result.scalars().all()
    if not images:
        raise HTTPException(status_code=400, detail="분석할 사진이 없어요")

    prompt = (
        "누수 진단 전문가로서 이미지를 분석하고 JSON으로만 답변하세요. "
        "필드: findings, work_items, materials, confidence, questions_for_user."
    )

    ai_response = await gemini_service.analyze_with_images(
        prompt=prompt,
        image_paths=[img.file_url for img in images],
        additional_context=payload.extra_context,
    )
    if ai_response.success and ai_response.data:
        normalized = _normalize_vision_result(ai_response.data)
        model_name = ai_response.model_used or "gemini-3.0-flash"
    else:
        normalized = _fallback_vision_result()
        model_name = "fallback-mock"

    confidence = Decimal(str(normalized.get("confidence", 0.0)))
    result_row = VisionResult(
        case_id=case_id,
        model=model_name,
        params_json={
            "temperature": 0.2,
            "response_mime_type": "application/json",
            "generated_at": datetime.utcnow().isoformat(),
        },
        result_json=normalized,
        confidence=confidence,
    )
    db.add(result_row)
    case.status = CaseStatus.VISION_READY
    case.updated_at = datetime.utcnow()
    await db.commit()
    await db.refresh(result_row)

    return APIResponse.ok(
        VisionResponse(
            id=result_row.id,
            case_id=result_row.case_id,
            model=result_row.model,
            result_json=result_row.result_json,
            confidence=float(result_row.confidence) if result_row.confidence is not None else None,
            created_at=result_row.created_at,
            updated_at=result_row.updated_at,
        )
    )


@router.patch("/cases/{case_id}/vision", response_model=APIResponse[VisionResponse])
async def update_case_vision(
    case_id: int,
    payload: VisionPatchRequest,
    db: DBSession,
    current_user: CurrentUser,
):
    case_result = await db.execute(select(Case).where(Case.id == case_id))
    case = case_result.scalar_one_or_none()
    if not case or not _can_access_case(case, current_user):
        raise HTTPException(status_code=404, detail="케이스를 찾을 수 없어요")

    vision_result = await db.execute(
        select(VisionResult)
        .where(VisionResult.case_id == case_id)
        .order_by(VisionResult.created_at.desc())
    )
    vision = vision_result.scalars().first()
    if not vision:
        raise HTTPException(status_code=404, detail="수정할 분석 결과가 없어요")

    normalized = _normalize_vision_result(payload.result_json)
    vision.result_json = normalized
    if payload.confidence is not None:
        vision.confidence = Decimal(str(max(0.0, min(payload.confidence, 1.0))))
    vision.updated_at = datetime.utcnow()
    case.updated_at = datetime.utcnow()
    await db.commit()
    await db.refresh(vision)

    return APIResponse.ok(
        VisionResponse(
            id=vision.id,
            case_id=vision.case_id,
            model=vision.model,
            result_json=vision.result_json,
            confidence=float(vision.confidence) if vision.confidence is not None else None,
            created_at=vision.created_at,
            updated_at=vision.updated_at,
        )
    )


@router.post("/cases/{case_id}/estimate", response_model=APIResponse[CaseEstimateResponse])
async def create_case_estimate(case_id: int, db: DBSession, current_user: CurrentUser):
    case_result = await db.execute(select(Case).where(Case.id == case_id))
    case = case_result.scalar_one_or_none()
    if not case or not _can_access_case(case, current_user):
        raise HTTPException(status_code=404, detail="케이스를 찾을 수 없어요")

    season_result = await db.execute(select(Season).where(Season.id == case.season_id))
    season = season_result.scalar_one_or_none()
    if not season or not season.is_active:
        raise HTTPException(status_code=400, detail="활성 시즌에서만 견적 생성할 수 있어요")

    vision_result = await db.execute(
        select(VisionResult)
        .where(VisionResult.case_id == case_id)
        .order_by(VisionResult.created_at.desc())
    )
    vision = vision_result.scalars().first()
    if not vision:
        raise HTTPException(status_code=400, detail="진단 결과가 없어요")

    await _ensure_default_estimation_category(db, case.season_id)
    category_result = await db.execute(
        select(SeasonCategory).where(
            SeasonCategory.season_id == case.season_id,
            SeasonCategory.purpose == SeasonCategoryPurpose.ESTIMATION,
            SeasonCategory.is_enabled == True,  # noqa: E712
        )
    )
    estimation_categories = list(category_result.scalars().all())
    estimation_category_names = {c.name for c in estimation_categories}
    if not estimation_category_names:
        raise HTTPException(status_code=400, detail="활성화된 적산 카테고리가 없어요")

    docs_result = await db.execute(
        select(SeasonDocument).where(
            SeasonDocument.season_id == case.season_id,
            SeasonDocument.status == DocumentStatus.DONE,
            SeasonDocument.is_enabled == True,  # noqa: E712
            SeasonDocument.category.in_(list(estimation_category_names)),
        )
    )
    estimation_docs = list(docs_result.scalars().all())
    estimation_doc_ids = {d.id for d in estimation_docs}
    if not estimation_doc_ids:
        raise HTTPException(status_code=400, detail="적산 자료 문서 인덱싱이 필요해요")

    cost_result = await db.execute(
        select(CostItem).where(
            CostItem.season_id == case.season_id,
            CostItem.source_doc_id.in_(list(estimation_doc_ids)),
        )
    )
    cost_items = cost_result.scalars().all()
    cost_items_lower = [(c, (c.item_name or "").lower(), (c.spec or "").lower(), (c.unit or "").lower()) for c in cost_items]

    payload = vision.result_json or {}
    candidates: list[dict] = []
    for work in payload.get("work_items", []) or []:
        if isinstance(work, dict):
            candidates.append({"kind": "work", **work})
    for material in payload.get("materials", []) or []:
        if isinstance(material, dict):
            candidates.append({"kind": "material", **material})

    lines: list[dict] = []
    matched_doc_ids: set[int] = set()
    for item in candidates:
        target_name = str(item.get("name", "")).strip()
        spec_hint = str(item.get("spec_hint", "") or item.get("spec", "")).strip().lower()
        unit_hint = str(item.get("unit_hint", "")).strip().lower()
        target_lower = target_name.lower()

        best_row: Optional[CostItem] = None
        best_score = -1
        for row, name_lower, spec_lower, unit_lower in cost_items_lower:
            score = 0
            if target_lower and target_lower in name_lower:
                score += 8
            if target_lower and name_lower in target_lower:
                score += 4
            if spec_hint and spec_hint in spec_lower:
                score += 3
            if unit_hint and unit_hint == unit_lower:
                score += 2
            if score > best_score:
                best_score = score
                best_row = row

        quantity = _extract_quantity(item)
        if best_row and best_score > 0:
            unit_price = float(best_row.unit_price)
            amount = round(quantity * unit_price, 2)
            matched_doc_ids.add(best_row.source_doc_id)
            evidence = [
                {
                    "doc_title": best_row.source_doc_title,
                    "season_name": season.name,
                    "page": best_row.source_page,
                    "table_id": best_row.table_id,
                    "row_id": best_row.row_id,
                    "row_text": best_row.row_text,
                }
            ]
            lines.append(
                {
                    "work": item.get("kind"),
                    "item_name": best_row.item_name,
                    "spec": best_row.spec,
                    "unit": best_row.unit,
                    "quantity": quantity,
                    "unit_price": unit_price,
                    "amount": amount,
                    "optional": best_score < 8,
                    "evidence": evidence,
                }
            )
        else:
            lines.append(
                {
                    "work": item.get("kind"),
                    "item_name": target_name or "미확인 항목",
                    "spec": item.get("spec_hint"),
                    "unit": item.get("unit_hint") or "ea",
                    "quantity": quantity,
                    "unit_price": 0.0,
                    "amount": 0.0,
                    "optional": True,
                    "evidence": [
                        {
                            "doc_title": "매칭 실패(사용자 확인 필요)",
                            "season_name": season.name,
                            "page": 0,
                            "table_id": None,
                            "row_id": None,
                            "row_text": "자동 매칭에 실패했습니다.",
                        }
                    ],
                }
            )

    subtotal = round(sum(float(line["amount"]) for line in lines), 2)
    vat = round(subtotal * 0.1, 2)
    total = round(subtotal + vat, 2)

    doc_hashes: list[str] = []
    if matched_doc_ids:
        doc_result = await db.execute(select(SeasonDocument).where(SeasonDocument.id.in_(matched_doc_ids)))
        docs = doc_result.scalars().all()
        doc_hashes = [d.version_hash for d in docs]
    version_hash_snapshot = ",".join(sorted(set(doc_hashes))) or "no-matched-doc"

    latest = await _latest_estimate(case.id, db)
    next_version = (latest.version + 1) if latest else 1
    estimate = CaseEstimate(
        case_id=case.id,
        version=next_version,
        items_json=lines,
        totals_json={
            "subtotal": subtotal,
            "vat_amount": vat,
            "total_amount": total,
        },
        version_hash_snapshot=version_hash_snapshot,
    )
    db.add(estimate)

    case.status = CaseStatus.ESTIMATED
    case.updated_at = datetime.utcnow()
    await db.commit()
    await db.refresh(estimate)

    return APIResponse.ok(
        CaseEstimateResponse(
            id=estimate.id,
            case_id=estimate.case_id,
            version=estimate.version,
            items=estimate.items_json,
            totals=estimate.totals_json,
            version_hash_snapshot=estimate.version_hash_snapshot,
            created_at=estimate.created_at,
        )
    )


@router.get("/cases/{case_id}/estimate", response_model=APIResponse[CaseEstimateResponse])
async def get_case_estimate(case_id: int, db: DBSession, current_user: CurrentUser):
    case_result = await db.execute(select(Case).where(Case.id == case_id))
    case = case_result.scalar_one_or_none()
    if not case or not _can_access_case(case, current_user):
        raise HTTPException(status_code=404, detail="케이스를 찾을 수 없어요")

    estimate = await _latest_estimate(case_id, db)
    if not estimate:
        raise HTTPException(status_code=404, detail="견적 결과가 없어요")

    return APIResponse.ok(
        CaseEstimateResponse(
            id=estimate.id,
            case_id=estimate.case_id,
            version=estimate.version,
            items=estimate.items_json,
            totals=estimate.totals_json,
            version_hash_snapshot=estimate.version_hash_snapshot,
            created_at=estimate.created_at,
        )
    )


@router.get("/cases/{case_id}/estimate.csv")
async def export_case_estimate_csv(case_id: int, db: DBSession, current_user: CurrentUser):
    case_result = await db.execute(select(Case).where(Case.id == case_id))
    case = case_result.scalar_one_or_none()
    if not case or not _can_access_case(case, current_user):
        raise HTTPException(status_code=404, detail="케이스를 찾을 수 없어요")

    estimate = await _latest_estimate(case_id, db)
    if not estimate:
        raise HTTPException(status_code=404, detail="견적 결과가 없어요")

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["공종", "품목명", "규격", "단위", "수량", "단가", "금액", "근거"])
    for line in estimate.items_json:
        evidence = (line.get("evidence") or [{}])[0]
        ref = (
            f"{evidence.get('doc_title', '')} / p.{evidence.get('page', '')} / "
            f"{evidence.get('table_id', '')} / {evidence.get('row_id', '')}"
        )
        writer.writerow(
            [
                line.get("work", ""),
                line.get("item_name", ""),
                line.get("spec", ""),
                line.get("unit", ""),
                line.get("quantity", 0),
                line.get("unit_price", 0),
                line.get("amount", 0),
                ref,
            ]
        )

    csv_bytes = output.getvalue().encode("utf-8-sig")
    file_name = f"case-{case_id}-estimate-v{estimate.version}.csv"
    file_url = await storage_service.save_bytes(
        data=csv_bytes,
        category="exports",
        subfolder=str(case_id),
        filename=file_name,
    )
    db.add(
        EstimateExport(
            estimate_id=estimate.id,
            file_type=ExportFileType.CSV,
            file_url=file_url,
        )
    )
    await db.commit()

    return StreamingResponse(
        io.BytesIO(csv_bytes),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{file_name}"'},
    )


@router.get("/cases/{case_id}/estimate.xlsx")
async def export_case_estimate_xlsx(case_id: int, db: DBSession, current_user: CurrentUser):
    case_result = await db.execute(select(Case).where(Case.id == case_id))
    case = case_result.scalar_one_or_none()
    if not case or not _can_access_case(case, current_user):
        raise HTTPException(status_code=404, detail="케이스를 찾을 수 없어요")

    estimate = await _latest_estimate(case_id, db)
    if not estimate:
        raise HTTPException(status_code=404, detail="견적 결과가 없어요")

    wb = Workbook()
    ws_estimate = wb.active
    ws_estimate.title = "견적서"
    ws_estimate.append(["공종", "품목명", "규격", "단위", "수량", "단가", "금액"])
    for line in estimate.items_json:
        ws_estimate.append(
            [
                line.get("work", ""),
                line.get("item_name", ""),
                line.get("spec", ""),
                line.get("unit", ""),
                line.get("quantity", 0),
                line.get("unit_price", 0),
                line.get("amount", 0),
            ]
        )
    ws_estimate.append([])
    ws_estimate.append(["", "", "", "", "", "공급가액", estimate.totals_json.get("subtotal", 0)])
    ws_estimate.append(["", "", "", "", "", "부가세", estimate.totals_json.get("vat_amount", 0)])
    ws_estimate.append(["", "", "", "", "", "합계", estimate.totals_json.get("total_amount", 0)])

    ws_evidence = wb.create_sheet("근거")
    ws_evidence.append(["품목명", "문서", "시즌", "페이지", "테이블", "행", "원문"])
    for line in estimate.items_json:
        for ev in line.get("evidence", []):
            ws_evidence.append(
                [
                    line.get("item_name", ""),
                    ev.get("doc_title", ""),
                    ev.get("season_name", ""),
                    ev.get("page", ""),
                    ev.get("table_id", ""),
                    ev.get("row_id", ""),
                    ev.get("row_text", ""),
                ]
            )

    output = io.BytesIO()
    wb.save(output)
    xlsx_bytes = output.getvalue()

    file_name = f"case-{case_id}-estimate-v{estimate.version}.xlsx"
    file_url = await storage_service.save_bytes(
        data=xlsx_bytes,
        category="exports",
        subfolder=str(case_id),
        filename=file_name,
    )
    db.add(
        EstimateExport(
            estimate_id=estimate.id,
            file_type=ExportFileType.XLSX,
            file_url=file_url,
        )
    )
    await db.commit()

    return StreamingResponse(
        io.BytesIO(xlsx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{file_name}"'},
    )
